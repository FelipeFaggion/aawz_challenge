from docx import Document
import re
import openpyxl

docx_path = '.../Partnership.docx' #Add your path to the docx file
doc = Document(docx_path)

excel_path = 'socios_info.xlsx'
workbook = openpyxl.Workbook()
sheet = workbook.active
sheet.title = "Socios Info"

sheet.append(['Nome', 'Cotas'])

for para in doc.paragraphs:
    if "portador" in para.text:
        lines = para.text.split(',')
        nome = lines[0].split(" ")[1] + " " + lines[0].split(" ")[2]
        cota = lines[-1].strip()

        match = re.search(r'\d+', cota) 
        if match:
            numero_cotas = int(match.group(0)) 

            sheet.append([nome, numero_cotas])

workbook.save(excel_path)
print(f"Dados salvos em {excel_path}")